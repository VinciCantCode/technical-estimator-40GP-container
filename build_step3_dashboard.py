"""
Step 3: Build the Executive Interactive Container Calculator Dashboard (1_Container_Calculator)
and assemble the complete 3-Tab Excel Model:
  Tab 1: 1_Container_Calculator (Interactive Dashboard & All-Schedule Comparison for Selected NPS)
  Tab 2: 2_Selectable_Master_Table (528-Row Filterable Master Table for CS & SS)
  Tab 3: 3_ASME_Spec_DB (ASME B36.10M & B36.19M Engineering Specification Database)
"""

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule
from pipe_data import ASME_PIPE_SPECS, SCHEDULES
from build_step2_master_table import build_step2_workbook


def build_complete_excel_model(filename="Precise_Pipe_40GP_Container_Model.xlsx"):
    # First build Step 1 (3_ASME_Spec_DB) and Step 2 (2_Selectable_Master_Table)
    build_step2_workbook(filename)
    wb = openpyxl.load_workbook(filename)

    # Create Tab 1: 1_Container_Calculator at index 0
    ws = wb.create_sheet(title="1_Container_Calculator", index=0)
    ws.views.sheetView[0].showGridLines = True

    # Color Palette (Precise Pipe Corporate Navy, Steel Slate, Emerald & Amber)
    NAVY_DARK = PatternFill(start_color="0F2C59", end_color="0F2C59", fill_type="solid")
    NAVY_SUB = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    SEC_HDR_FILL = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    SUB_HDR_FILL = PatternFill(start_color="334155", end_color="334155", fill_type="solid")
    EMERALD_HDR = PatternFill(start_color="047857", end_color="047857", fill_type="solid")
    INDIGO_HDR = PatternFill(start_color="3730A3", end_color="3730A3", fill_type="solid")

    INPUT_FILL = PatternFill(start_color="FFFBEB", end_color="FFFBEB", fill_type="solid")     # Warm Amber input
    LABEL_FILL = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")     # Crisp Slate label
    VALUE_FILL = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")     # Pure White value
    ZEBRA_FILL = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    KPI_GREEN_BG = PatternFill(start_color="ECFDF5", end_color="ECFDF5", fill_type="solid")   # Emerald KPI card
    KPI_BLUE_BG = PatternFill(start_color="EFF6FF", end_color="EFF6FF", fill_type="solid")    # Royal Blue KPI card
    KPI_AMBER_BG = PatternFill(start_color="FFFBEB", end_color="FFFBEB", fill_type="solid")   # Commercial KPI card
    INSIGHT_BG = PatternFill(start_color="F0FDFA", end_color="F0FDFA", fill_type="solid")     # Teal commercial banner

    # Fonts
    FONT_TITLE = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
    FONT_SUBTITLE = Font(name="Calibri", size=10, italic=True, color="E2E8F0")
    FONT_SEC_HDR = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    FONT_LBL = Font(name="Calibri", size=10, bold=True, color="1E293B")
    FONT_INPUT = Font(name="Calibri", size=11, bold=True, color="B45309")
    FONT_VAL_BOLD = Font(name="Calibri", size=10, bold=True, color="0F2C59")
    FONT_VAL_REG = Font(name="Calibri", size=10, bold=False, color="1E293B")
    FONT_KPI_BIG_GRN = Font(name="Calibri", size=22, bold=True, color="047857")
    FONT_KPI_BIG_BLU = Font(name="Calibri", size=22, bold=True, color="1D4ED8")
    FONT_KPI_BIG_AMB = Font(name="Calibri", size=15, bold=True, color="B45309")
    FONT_KPI_SUB = Font(name="Calibri", size=9, italic=True, color="475569")
    FONT_INSIGHT = Font(name="Calibri", size=10, bold=True, color="0F766E")

    # Borders
    THIN_GRAY = Side(border_style="thin", color="CBD5E1")
    MED_NAVY = Side(border_style="medium", color="0F2C59")
    MED_EMERALD = Side(border_style="medium", color="047857")
    MED_AMBER = Side(border_style="medium", color="D97706")
    THIN_AMBER = Side(border_style="thin", color="F59E0B")

    BORDER_STD = Border(left=THIN_GRAY, right=THIN_GRAY, top=THIN_GRAY, bottom=THIN_GRAY)
    BORDER_INPUT = Border(left=THIN_AMBER, right=THIN_AMBER, top=THIN_AMBER, bottom=THIN_AMBER)
    BORDER_HDR = Border(left=THIN_GRAY, right=THIN_GRAY, top=MED_NAVY, bottom=MED_NAVY)
    BORDER_KPI_GRN = Border(left=MED_EMERALD, right=MED_EMERALD, top=MED_EMERALD, bottom=MED_EMERALD)
    BORDER_KPI_BLU = Border(left=MED_NAVY, right=MED_NAVY, top=MED_NAVY, bottom=MED_NAVY)
    BORDER_KPI_AMB = Border(left=MED_AMBER, right=MED_AMBER, top=MED_AMBER, bottom=MED_AMBER)

    def style_merged_range(ws_obj, cell_range, value=None, font=None, fill=None, border=None, alignment=None, num_format=None):
        ws_obj.merge_cells(cell_range)
        start_col, start_row, end_col, end_row = openpyxl.utils.range_boundaries(cell_range)
        top_left = ws_obj.cell(row=start_row, column=start_col)
        if value is not None:
            top_left.value = value
        if font is not None:
            top_left.font = font
        if alignment is not None:
            top_left.alignment = alignment
        if num_format is not None:
            top_left.number_format = num_format
        for r in range(start_row, end_row + 1):
            for c in range(start_col, end_col + 1):
                cell = ws_obj.cell(row=r, column=c)
                if fill is not None:
                    cell.fill = fill
                if border is not None:
                    cell.border = border

    # 1. Executive Title Banner (Rows 1 to 2)
    style_merged_range(
        ws, "A1:N1",
        "PRECISE PIPE — 40GP CONTAINER LOADING & COMMERCIAL ESTIMATOR MODEL",
        font=FONT_TITLE, fill=NAVY_DARK,
        alignment=Alignment(horizontal="left", vertical="center", indent=1)
    )
    ws.row_dimensions[1].height = 32

    style_merged_range(
        ws, "A2:N2",
        "Interactive Calculator & Schedule Matrix | Select Pipe Size (NPS), Schedule & Material in Yellow Cells (Col D) | ASME B36.10M (CS) & B36.19M (SS)",
        font=FONT_SUBTITLE, fill=NAVY_SUB,
        alignment=Alignment(horizontal="left", vertical="center", indent=1)
    )
    ws.row_dimensions[2].height = 20
    ws.row_dimensions[3].height = 10

    # =========================================================================
    # LEFT PANEL (Cols A to E, Rows 4 to 18): USER INPUTS & CONFIGURATION
    # =========================================================================
    style_merged_range(
        ws, "A4:E4",
        "1. PIPE SPECIFICATION INPUTS (SELECT DROPDOWNS)",
        font=FONT_SEC_HDR, fill=SEC_HDR_FILL, border=BORDER_HDR,
        alignment=Alignment(horizontal="left", vertical="center", indent=1)
    )
    ws.row_dimensions[4].height = 24

    pipe_inputs = [
        (5, "Material Grade [Select]", "Carbon Steel (CS)", "@", True),
        (6, "Nominal Pipe Size - NPS [Select]", "6''", "@", True),
        (7, "Pipe Schedule [Select]", "40", "@", True),
        (8, "Single Pipe Length (m) [Select]", 11.80, "0.00", True),
        (9, "ASME Spec Status", '=IF(ISNUMBER(I12), "VALID ASME SPEC", "SCHEDULE N/A FOR THIS NPS (SEE TABLE BELOW)")', None, False),
    ]

    for r_idx, lbl_text, default_val, num_fmt, is_input in pipe_inputs:
        ws.row_dimensions[r_idx].height = 22
        style_merged_range(
            ws, f"A{r_idx}:C{r_idx}", lbl_text,
            font=FONT_LBL, fill=LABEL_FILL, border=BORDER_STD,
            alignment=Alignment(horizontal="left", vertical="center", indent=1)
        )
        style_merged_range(
            ws, f"D{r_idx}:E{r_idx}", default_val,
            font=FONT_INPUT if is_input else FONT_VAL_BOLD,
            fill=INPUT_FILL if is_input else KPI_BLUE_BG,
            border=BORDER_INPUT if is_input else BORDER_STD,
            alignment=Alignment(horizontal="center", vertical="center"),
            num_format=num_fmt
        )

    # Section 2 on Left Panel: 40GP Container & Commercial Parameters (Rows 11 to 18)
    ws.row_dimensions[10].height = 8
    style_merged_range(
        ws, "A11:E11",
        "2. 40GP CONTAINER & COMMERCIAL PARAMETERS",
        font=FONT_SEC_HDR, fill=SEC_HDR_FILL, border=BORDER_HDR,
        alignment=Alignment(horizontal="left", vertical="center", indent=1)
    )
    ws.row_dimensions[11].height = 24

    container_inputs = [
        (12, "40GP Internal Length (m)", 12.032, "0.000", False),
        (13, "Effective Loading Width (mm) [Input]", 2300, "#,##0", True),
        (14, "Effective Loading Height (mm) [Input]", 2280, "#,##0", True),
        (15, "Max Container Payload (MT) [Input]", 26.00, "0.00", True),
        (16, "Cross-Section Packing Mode [Select]", "Staggered (Hexagonal)", None, True),
        (17, "Est. Ocean Freight / 40GP (USD) [Input]", 3500, "$#,##0", True),
        (18, "Nested Pipe Clearance Margin (mm)", 25, "#,##0", True),
    ]

    for r_idx, lbl_text, default_val, num_fmt, is_input in container_inputs:
        ws.row_dimensions[r_idx].height = 21
        style_merged_range(
            ws, f"A{r_idx}:C{r_idx}", lbl_text,
            font=FONT_LBL, fill=LABEL_FILL, border=BORDER_STD,
            alignment=Alignment(horizontal="left", vertical="center", indent=1)
        )
        style_merged_range(
            ws, f"D{r_idx}:E{r_idx}", default_val,
            font=FONT_INPUT if is_input else FONT_VAL_REG,
            fill=INPUT_FILL if is_input else VALUE_FILL,
            border=BORDER_INPUT if is_input else BORDER_STD,
            alignment=Alignment(horizontal="center", vertical="center"),
            num_format=num_fmt
        )

    # =========================================================================
    # RIGHT PANEL TOP (Cols G to N, Rows 4 to 8): EXECUTIVE KPI OUTPUT CARDS
    # =========================================================================
    # Card 1: 40GP ACTUAL PIECES (Cols G to I)
    style_merged_range(
        ws, "G4:I4", "OUTPUT 1: 40GP TOTAL PIECES",
        font=FONT_SEC_HDR, fill=EMERALD_HDR, border=BORDER_HDR,
        alignment=Alignment(horizontal="center", vertical="center")
    )
    style_merged_range(
        ws, "G5:I7", '=IF(ISNUMBER(I12), M16, 0)',
        font=FONT_KPI_BIG_GRN, fill=KPI_GREEN_BG, border=BORDER_KPI_GRN,
        alignment=Alignment(horizontal="center", vertical="center"),
        num_format='#,##0" pcs"'
    )
    style_merged_range(
        ws, "G8:I8", '="Total Length: " & TEXT(G5*D8, "#,##0.0") & " m (" & TEXT(G5*D8*3.28084, "#,##0") & " ft)"',
        font=FONT_KPI_SUB, fill=KPI_GREEN_BG, border=BORDER_STD,
        alignment=Alignment(horizontal="center", vertical="center")
    )

    # Card 2: 40GP TOTAL METRIC TONS (Cols J to L)
    style_merged_range(
        ws, "J4:L4", "OUTPUT 2: 40GP METRIC TONS (MT)",
        font=FONT_SEC_HDR, fill=NAVY_SUB, border=BORDER_HDR,
        alignment=Alignment(horizontal="center", vertical="center")
    )
    style_merged_range(
        ws, "J5:L7", '=IF(ISNUMBER(I12), M17, 0)',
        font=FONT_KPI_BIG_BLU, fill=KPI_BLUE_BG, border=BORDER_KPI_BLU,
        alignment=Alignment(horizontal="center", vertical="center"),
        num_format='#,##0.00" MT"'
    )
    style_merged_range(
        ws, "J8:L8", '="Net Weight: " & TEXT(J5*1000, "#,##0") & " kg (" & TEXT(J5*2204.62, "#,##0") & " lbs)"',
        font=FONT_KPI_SUB, fill=KPI_BLUE_BG, border=BORDER_STD,
        alignment=Alignment(horizontal="center", vertical="center")
    )

    # Card 3: BOTTLENECK & FREIGHT ECONOMICS (Cols M to N)
    style_merged_range(
        ws, "M4:N4", "BOTTLENECK & FREIGHT",
        font=FONT_SEC_HDR, fill=INDIGO_HDR, border=BORDER_HDR,
        alignment=Alignment(horizontal="center", vertical="center")
    )
    style_merged_range(
        ws, "M5:N6", '=IF(G5=0, "N/A", IF(M15<=M14, "WEIGHT LIMITED", "VOLUME LIMITED"))',
        font=FONT_KPI_BIG_AMB, fill=KPI_AMBER_BG, border=BORDER_KPI_AMB,
        alignment=Alignment(horizontal="center", vertical="center")
    )
    style_merged_range(
        ws, "M7:N7", '=IF(J5>0, "Freight: $" & TEXT(D17/J5, "#,##0.00") & " / MT", "Freight: N/A")',
        font=FONT_VAL_BOLD, fill=KPI_AMBER_BG, border=BORDER_STD,
        alignment=Alignment(horizontal="center", vertical="center")
    )
    style_merged_range(
        ws, "M8:N8", '=IF(G5>0, "Unit Freight: $" & TEXT(D17/(G5*D8), "0.00") & " / m", "Unit Freight: N/A")',
        font=FONT_KPI_SUB, fill=KPI_AMBER_BG, border=BORDER_STD,
        alignment=Alignment(horizontal="center", vertical="center")
    )

    # =========================================================================
    # RIGHT PANEL MIDDLE (Cols G to N, Rows 10 to 18): ENGINEERING & PACKING BREAKDOWN
    # =========================================================================
    style_merged_range(
        ws, "G10:J10", "3. SELECTED PIPE PHYSICAL SPECS (ASME)",
        font=FONT_SEC_HDR, fill=SUB_HDR_FILL, border=BORDER_HDR,
        alignment=Alignment(horizontal="left", vertical="center", indent=1)
    )
    style_merged_range(
        ws, "K10:N10", "4. 40GP CONTAINER PACKING GEOMETRY",
        font=FONT_SEC_HDR, fill=SUB_HDR_FILL, border=BORDER_HDR,
        alignment=Alignment(horizontal="left", vertical="center", indent=1)
    )
    ws.row_dimensions[10].height = 22

    # Formulas referencing 3_ASME_Spec_DB (Rows 12 to 47 in 3_ASME_Spec_DB)
    # Col A = NPS, Col B = DN, Col C = OD(in), Col D = OD(mm), Col E..V = Sch 5S..XXS
    eng_specs_rows = [
        (11, "Outer Diameter - OD (mm)", '=IFERROR(INDEX(\'3_ASME_Spec_DB\'!$D$12:$D$47, MATCH(D6, \'3_ASME_Spec_DB\'!$A$12:$A$47, 0)), 0)', "0.00"),
        (12, "Wall Thickness - WT (mm)", '=IFERROR(INDEX(\'3_ASME_Spec_DB\'!$E$12:$V$47, MATCH(D6, \'3_ASME_Spec_DB\'!$A$12:$A$47, 0), MATCH("Sch "&D7, \'3_ASME_Spec_DB\'!$E$11:$V$11, 0)), "N/A")', "0.00"),
        (13, "Outer / Wall (Inches)", '=TEXT(I11/25.4,"0.000") & " in OD x " & IF(ISNUMBER(I12), TEXT(I12/25.4,"0.000") & " in WT", "N/A")', None),
        (14, "Inner Diameter - ID (mm)", '=IF(ISNUMBER(I12), ROUND(I11 - 2*I12, 2), "N/A")', "0.00"),
        (15, "Material Weight Factor C", '=IF(D5="Carbon Steel (CS)", \'3_ASME_Spec_DB\'!D6, IF(D5="Stainless Steel 304 (SS304)", \'3_ASME_Spec_DB\'!D7, \'3_ASME_Spec_DB\'!D8))', "0.0000000"),
        (16, "Linear Unit Weight (kg/m)", '=IF(ISNUMBER(I12), ROUND((I11 - I12) * I12 * I15, 3), 0)', "#,##0.000"),
        (17, "Linear Unit Weight (lb/ft)", '=IF(ISNUMBER(I12), ROUND(I16 * 0.671969, 3), 0)', "#,##0.000"),
        (18, "Single Pipe Weight (kg/pc)", '=IF(ISNUMBER(I12), ROUND(I16 * D8, 2), 0)', "#,##0.00"),
    ]

    packing_geo_rows = [
        (11, "Lengthwise Tiers in 40GP", '=INT(D12 / D8)', "#,##0"),
        (12, "Bottom Row Fit (pcs / row)", '=IF(I11>0, INT(D13 / I11), 0)', "#,##0"),
        (13, "Vertical Stacked Layers", '=IF(M12=0, 0, IF(OR(D16="Square Grid", M12=1), INT(D14/I11), IF(D14<I11, 0, 1+INT((D14-I11)/(0.8660254*I11)))))', "#,##0"),
        (14, "Max Pieces by Volume (pcs)", '=IF(OR(M12=0,M13=0), 0, IF(OR(D16="Square Grid", M12=1), M12*M13*M11, (INT((M13+1)/2)*M12 + INT(M13/2)*IF((D13-M12*I11)>=0.5*I11, M12, M12-1))*M11))', "#,##0"),
        (15, "Max Pieces by Weight (pcs)", '=IF(I18>0, INT((D15*1000) / I18), 0)', "#,##0"),
        (16, "40GP Actual Loaded Pieces", '=IF(ISNUMBER(I12), MIN(M14, M15), 0)', "#,##0"),
        (17, "40GP Actual Net Weight (MT)", '=ROUND((M16 * I18) / 1000, 2)', "#,##0.00"),
        (18, "Weight Util. / Space Slot Util.", '=TEXT(IF(D15>0, M17/D15, 0), "0.0%") & " Wt  |  " & TEXT(IF(M14>0, M16/M14, 0), "0.0%") & " Space"', None),
    ]

    for (r_idx, lbl_l, form_l, fmt_l), (_, lbl_r, form_r, fmt_r) in zip(eng_specs_rows, packing_geo_rows):
        ws.row_dimensions[r_idx].height = 21
        # Left spec item (G:H label, I:J value)
        style_merged_range(
            ws, f"G{r_idx}:H{r_idx}", lbl_l,
            font=FONT_LBL, fill=LABEL_FILL, border=BORDER_STD,
            alignment=Alignment(horizontal="left", vertical="center", indent=1)
        )
        style_merged_range(
            ws, f"I{r_idx}:J{r_idx}", form_l,
            font=FONT_VAL_BOLD if r_idx in (11, 12, 16, 18) else FONT_VAL_REG,
            fill=VALUE_FILL, border=BORDER_STD,
            alignment=Alignment(horizontal="right" if fmt_l else "center", vertical="center"),
            num_format=fmt_l
        )
        # Right packing item (K:L label, M:N value)
        style_merged_range(
            ws, f"K{r_idx}:L{r_idx}", lbl_r,
            font=FONT_LBL, fill=LABEL_FILL, border=BORDER_STD,
            alignment=Alignment(horizontal="left", vertical="center", indent=1)
        )
        is_key_out = r_idx in (16, 17)
        style_merged_range(
            ws, f"M{r_idx}:N{r_idx}", form_r,
            font=FONT_VAL_BOLD if is_key_out else FONT_VAL_REG,
            fill=KPI_GREEN_BG if is_key_out else VALUE_FILL,
            border=BORDER_STD,
            alignment=Alignment(horizontal="right" if fmt_r else "center", vertical="center"),
            num_format=fmt_r
        )

    # =========================================================================
    # ROW 20: AI & COMMERCIAL ANALYST INSIGHT BANNER (Telescoping / Optimization)
    # =========================================================================
    ws.row_dimensions[19].height = 8
    ws.row_dimensions[20].height = 26
    insight_formula = (
        '=IF(NOT(ISNUMBER(I12)), "COMMERCIAL ADVISOR: Selected Schedule (" & D7 & ") is not standard for NPS " & D6 & ". Review the live table below for all valid Schedules for NPS " & D6 & ".", '
        'IF(M15<=M14, "COMMERCIAL ADVISOR [WEIGHT LIMITED]: Container reaches Max Payload (" & TEXT(M17,"0.00") & " MT) at " & TEXT(M16/M14,"0.0%") & " space occupancy. Increasing container weight limit or using lighter wall schedule increases piece count per box.", '
        'IF(I14>80, "COMMERCIAL ADVISOR [VOLUME LIMITED — PIPE NESTING OPPORTUNITY]: Space is 100% full at only " & TEXT(M17,"0.00") & " MT (" & TEXT(D15-M17,"0.00") & " MT unused payload!). Inner Diameter is " & TEXT(I14,"0.0") & " mm -> You can telescope/nest smaller pipes (OD <= " & TEXT(I14-D18,"0") & " mm) inside each pipe to ship extra tonnage for FREE ocean freight!", '
        '"COMMERCIAL ADVISOR [VOLUME LIMITED]: Container space is 100% utilized (" & TEXT(M16,"#,##0") & " pcs / " & TEXT(M17,"0.00") & " MT).")))'
    )
    style_merged_range(
        ws, "A20:N20", insight_formula,
        font=FONT_INSIGHT, fill=INSIGHT_BG, border=BORDER_KPI_GRN,
        alignment=Alignment(horizontal="left", vertical="center", indent=1)
    )

    # =========================================================================
    # BOTTOM SECTION (Rows 22 to 41): ALL SCHEDULES COMPARISON TABLE FOR SELECTED NPS
    # =========================================================================
    ws.row_dimensions[21].height = 8
    style_merged_range(
        ws, "A22:N22",
        '="5. ALL PIPE SCHEDULES COMPARISON MATRIX FOR SELECTED SIZE: NPS " & D6 & " (" & D5 & " | Pipe Length: " & TEXT(D8,"0.00") & "m)"',
        font=FONT_SEC_HDR, fill=NAVY_DARK, border=BORDER_HDR,
        alignment=Alignment(horizontal="left", vertical="center", indent=1)
    )
    ws.row_dimensions[22].height = 24

    comp_headers = [
        ("Schedule", NAVY_DARK),
        ("Spec Status", NAVY_DARK),
        ("OD (mm)", SUB_HDR_FILL),
        ("WT (in)", SUB_HDR_FILL),
        ("WT (mm)", SUB_HDR_FILL),
        ("ID (mm)", SUB_HDR_FILL),
        ("Unit Wt\n(kg/m)", SUB_HDR_FILL),
        ("Piece Wt\n(kg/pc)", SUB_HDR_FILL),
        ("Max Pieces\n(Volume)", SUB_HDR_FILL),
        ("Max Pieces\n(Weight)", SUB_HDR_FILL),
        ("40GP Actual\nPieces (pcs)", EMERALD_HDR),
        ("40GP Actual\nMetric Tons (MT)", EMERALD_HDR),
        ("Payload\nUtil. (%)", INDIGO_HDR),
        ("Limiting Factor\n(Bottleneck)", INDIGO_HDR),
    ]

    ws.row_dimensions[23].height = 28
    for c_i, (h_title, h_fill) in enumerate(comp_headers, start=1):
        c = ws.cell(row=23, column=c_i, value=h_title)
        c.font = FONT_SEC_HDR
        c.fill = h_fill
        c.border = BORDER_HDR
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    start_comp_row = 24
    for idx, sch_code in enumerate(SCHEDULES):
        r = start_comp_row + idx
        ws.row_dimensions[r].height = 19
        row_fill = ZEBRA_FILL if idx % 2 == 1 else VALUE_FILL
        db_col_letter = get_column_letter(5 + idx)  # Col E is index 5 in 3_ASME_Spec_DB

        # Formulas for each schedule row r for the selected NPS ($D$6):
        f_wt_mm = f'=IFERROR(INDEX(\'3_ASME_Spec_DB\'!${db_col_letter}$12:${db_col_letter}$47, MATCH($D$6, \'3_ASME_Spec_DB\'!$A$12:$A$47, 0)), "-")'
        f_status = f'=IF(ISNUMBER(E{r}), "Available", "N/A")'
        f_od_mm = f'=IF(ISNUMBER(E{r}), $I$11, "-")'
        f_wt_in = f'=IF(ISNUMBER(E{r}), ROUND(E{r}/25.4, 3), "-")'
        f_id_mm = f'=IF(ISNUMBER(E{r}), ROUND(C{r}-2*E{r}, 2), "-")'
        f_unit_wt = f'=IF(ISNUMBER(E{r}), ROUND((C{r}-E{r})*E{r}*$I$15, 3), "-")'
        f_pc_wt = f'=IF(ISNUMBER(E{r}), ROUND(G{r}*$D$8, 2), "-")'
        f_max_vol = f'=IF(ISNUMBER(E{r}), $M$14, "-")'
        f_max_wt = f'=IF(ISNUMBER(E{r}), INT(($D$15*1000)/H{r}), "-")'
        f_act_pcs = f'=IF(ISNUMBER(E{r}), MIN(I{r}, J{r}), "-")'
        f_act_mt = f'=IF(ISNUMBER(E{r}), ROUND((K{r}*H{r})/1000, 2), "-")'
        f_wt_util = f'=IF(ISNUMBER(E{r}), L{r}/$D$15, "-")'
        f_bottleneck = f'=IF(NOT(ISNUMBER(E{r})), "-", IF(J{r}<=I{r}, "Weight Limited", "Volume Limited"))'

        row_vals = [
            sch_code,      # A: Schedule
            f_status,      # B: Spec Status
            f_od_mm,       # C: OD (mm)
            f_wt_in,       # D: WT (in)
            f_wt_mm,       # E: WT (mm)
            f_id_mm,       # F: ID (mm)
            f_unit_wt,     # G: Unit Wt (kg/m)
            f_pc_wt,       # H: Piece Wt (kg/pc)
            f_max_vol,     # I: Max Pieces (Volume)
            f_max_wt,      # J: Max Pieces (Weight)
            f_act_pcs,     # K: 40GP Actual Pieces
            f_act_mt,      # L: 40GP Actual MT
            f_wt_util,     # M: Payload Util (%)
            f_bottleneck,  # N: Limiting Factor
        ]

        for c_i, val in enumerate(row_vals, start=1):
            cell = ws.cell(row=r, column=c_i, value=val)
            cell.border = BORDER_STD
            if c_i in (11, 12):
                cell.fill = KPI_GREEN_BG
                cell.font = FONT_VAL_BOLD
            else:
                cell.fill = row_fill
                cell.font = FONT_VAL_BOLD if c_i in (1, 2, 14) else FONT_VAL_REG

            # Alignment
            if c_i in (1, 2, 14):
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="right", vertical="center")

            # Number Formats
            if c_i in (3, 5, 6, 8, 12):
                cell.number_format = "#,##0.00"
            elif c_i in (4, 7):
                cell.number_format = "#,##0.000"
            elif c_i in (9, 10, 11):
                cell.number_format = "#,##0"
            elif c_i == 13:
                cell.number_format = "0.0%"

    last_comp_row = start_comp_row + len(SCHEDULES) - 1

    # Conditional formatting on Comparison Table
    avail_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
    avail_font = Font(name="Calibri", size=10, bold=True, color="15803D")
    na_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    na_font = Font(name="Calibri", size=9, italic=True, color="94A3B8")

    ws.conditional_formatting.add(
        f"B24:B{last_comp_row}",
        CellIsRule(operator="equal", formula=['"Available"'], fill=avail_fill, font=avail_font)
    )
    ws.conditional_formatting.add(
        f"B24:B{last_comp_row}",
        CellIsRule(operator="equal", formula=['"N/A"'], fill=na_fill, font=na_font)
    )

    wt_lim_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
    wt_lim_font = Font(name="Calibri", size=10, bold=True, color="B45309")
    vol_lim_fill = PatternFill(start_color="DBEAFE", end_color="DBEAFE", fill_type="solid")
    vol_lim_font = Font(name="Calibri", size=10, bold=True, color="1D4ED8")

    ws.conditional_formatting.add(
        f"N24:N{last_comp_row}",
        CellIsRule(operator="equal", formula=['"Weight Limited"'], fill=wt_lim_fill, font=wt_lim_font)
    )
    ws.conditional_formatting.add(
        f"N24:N{last_comp_row}",
        CellIsRule(operator="equal", formula=['"Volume Limited"'], fill=vol_lim_fill, font=vol_lim_font)
    )

    # Data Validations for User Inputs on 1_Container_Calculator
    # All dropdowns use direct literal lists and full merged ranges for 100% compatibility
    dv_mat = DataValidation(
        type="list",
        formula1='"Carbon Steel (CS),Stainless Steel 304 (SS304),Stainless Steel 316 (SS316)"',
        allow_blank=False
    )
    ws.add_data_validation(dv_mat)
    dv_mat.add("D5:E5")

    nps_literal_list = '"' + ",".join([spec[0] for spec in ASME_PIPE_SPECS]) + '"'
    dv_nps = DataValidation(
        type="list",
        formula1=nps_literal_list,
        allow_blank=False
    )
    ws.add_data_validation(dv_nps)
    dv_nps.add("D6:E6")

    sch_literal_list = '"' + ",".join(SCHEDULES) + '"'
    dv_sch = DataValidation(
        type="list",
        formula1=sch_literal_list,
        allow_blank=False
    )
    ws.add_data_validation(dv_sch)
    dv_sch.add("D7:E7")

    dv_len = DataValidation(
        type="list",
        formula1='"5.80,6.00,11.80,12.00"',
        allow_blank=False
    )
    ws.add_data_validation(dv_len)
    dv_len.add("D8:E8")

    dv_pack = DataValidation(
        type="list",
        formula1='"Staggered (Hexagonal),Square Grid"',
        allow_blank=False
    )
    ws.add_data_validation(dv_pack)
    dv_pack.add("D16:E16")

    # Set clean column widths across A to N
    col_widths = {
        1: 12,  # A
        2: 13,  # B
        3: 11,  # C
        4: 13,  # D
        5: 13,  # E
        6: 4,   # F (Spacer / ID col in bottom table)
        7: 13,  # G
        8: 13,  # H
        9: 13,  # I
        10: 13, # J
        11: 15, # K
        12: 16, # L
        13: 13, # M
        14: 16, # N
    }
    # Make Col F wide enough for ID (mm) in bottom comparison table
    col_widths[6] = 11
    for c_i, w in col_widths.items():
        ws.column_dimensions[get_column_letter(c_i)].width = w

    # Set active sheet to 1_Container_Calculator so it opens first
    wb.active = 0
    wb.save(filename)

    # Validate all formulas across all worksheets
    errors = []
    formula_count = 0
    for sheet in wb.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    formula_count += 1
                    if cell.value.count("(") != cell.value.count(")") or cell.value.count('"') % 2 != 0:
                        errors.append((sheet.title, cell.coordinate, cell.value))

    print(f"Successfully assembled complete 3-Tab Excel Model in {filename}!")
    print(f"Validated {formula_count} dynamic Excel formulas across {len(wb.worksheets)} sheets. Syntax errors found: {len(errors)}")
    if errors:
        for err in errors:
            print("  ERROR:", err)


if __name__ == "__main__":
    build_complete_excel_model()
