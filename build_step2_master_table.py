"""
Step 2: Build the Selectable Master Table Sheet (2_Selectable_Master_Table)
alongside the Engineering Specification Database Sheet (3_ASME_Spec_DB)
for the Precise Pipe 40GP Container Loading Model.
"""

import math
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule
from pipe_data import ASME_PIPE_SPECS, SCHEDULES, CARBON_SCHEDULES, STAINLESS_SCHEDULES
from build_step1_db import create_step1_workbook, INCH_TO_MM


def build_step2_workbook(filename="Precise_Pipe_40GP_Container_Model.xlsx"):
    # First generate Step 1 workbook (3_ASME_Spec_DB)
    create_step1_workbook(filename)
    wb = openpyxl.load_workbook(filename)

    # Create 2_Selectable_Master_Table as the first tab (before 3_ASME_Spec_DB)
    ws = wb.create_sheet(title="2_Selectable_Master_Table", index=0)
    ws.views.sheetView[0].showGridLines = True

    # Styling definitions
    NAVY_HEADER = PatternFill(start_color="0F2C59", end_color="0F2C59", fill_type="solid")
    SLATE_SUBHEADER = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    PARAM_HDR_FILL = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    PARAM_LBL_FILL = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    INPUT_CELL_FILL = PatternFill(start_color="FFFBEB", end_color="FFFBEB", fill_type="solid")  # Soft amber for editable inputs
    CALC_CELL_FILL = PatternFill(start_color="EFF6FF", end_color="EFF6FF", fill_type="solid")   # Soft blue for auto-calculated params
    TBL_HDR_MAIN = PatternFill(start_color="0F2C59", end_color="0F2C59", fill_type="solid")
    TBL_HDR_CALC = PatternFill(start_color="334155", end_color="334155", fill_type="solid")
    TBL_HDR_OUT = PatternFill(start_color="047857", end_color="047857", fill_type="solid")      # Emerald green for key outputs

    ZEBRA_FILL = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    WHITE_FILL = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    OUTPUT_COL_FILL = PatternFill(start_color="ECFDF5", end_color="ECFDF5", fill_type="solid")  # Soft green tint for output columns

    FONT_TITLE = Font(name="Calibri", size=15, bold=True, color="FFFFFF")
    FONT_SUBTITLE = Font(name="Calibri", size=10, italic=True, color="E2E8F0")
    FONT_PARAM_HDR = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    FONT_PARAM_LBL = Font(name="Calibri", size=10, bold=True, color="1E293B")
    FONT_PARAM_INP = Font(name="Calibri", size=10, bold=True, color="B45309")
    FONT_PARAM_VAL = Font(name="Calibri", size=10, bold=True, color="0F2C59")
    FONT_TBL_HDR = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    FONT_REGULAR = Font(name="Calibri", size=10, bold=False, color="1E293B")
    FONT_BOLD = Font(name="Calibri", size=10, bold=True, color="0F172A")
    FONT_OUTPUT = Font(name="Calibri", size=10, bold=True, color="065F46")
    FONT_CS_TAG = Font(name="Calibri", size=10, bold=True, color="1E3A8A")
    FONT_SS_TAG = Font(name="Calibri", size=10, bold=True, color="0F766E")

    THIN_BORDER_SIDE = Side(border_style="thin", color="CBD5E1")
    MED_NAVY_SIDE = Side(border_style="medium", color="0F2C59")
    INPUT_BORDER_SIDE = Side(border_style="thin", color="F59E0B")

    BORDER_DATA = Border(left=THIN_BORDER_SIDE, right=THIN_BORDER_SIDE, top=THIN_BORDER_SIDE, bottom=THIN_BORDER_SIDE)
    BORDER_INPUT = Border(left=INPUT_BORDER_SIDE, right=INPUT_BORDER_SIDE, top=INPUT_BORDER_SIDE, bottom=INPUT_BORDER_SIDE)
    BORDER_HDR = Border(left=THIN_BORDER_SIDE, right=THIN_BORDER_SIDE, top=MED_NAVY_SIDE, bottom=MED_NAVY_SIDE)

    # 1. Title Banner
    ws.merge_cells("A1:U1")
    ws["A1"] = "PRECISE PIPE — 40GP CONTAINER LOADING MASTER SELECTABLE TABLE (ALL NPS & SCHEDULES)"
    ws["A1"].font = FONT_TITLE
    ws["A1"].fill = NAVY_HEADER
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[1].height = 30

    ws.merge_cells("A2:U2")
    ws["A2"] = "Use column filter dropdowns (Row 10) to filter by Material (Carbon Steel / Stainless Steel), Nominal Pipe Size (NPS 1/8\" to 48\"), Pipe Schedule, or Bottleneck Type. Yellow cells (Rows 5-8) are interactive parameters."
    ws["A2"].font = FONT_SUBTITLE
    ws["A2"].fill = SLATE_SUBHEADER
    ws["A2"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[2].height = 20

    # 2. Container & Packing Assumptions Control Block (Rows 4 to 8)
    ws.merge_cells("A4:D4")
    ws["A4"] = "40GP CONTAINER DIMENSIONS & WEIGHT LIMIT"
    ws["A4"].font = FONT_PARAM_HDR
    ws["A4"].fill = PARAM_HDR_FILL
    ws["A4"].alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells("E4:H4")
    ws["E4"] = "EFFECTIVE LOADING SPACE & PIPE LENGTH"
    ws["E4"].font = FONT_PARAM_HDR
    ws["E4"].fill = PARAM_HDR_FILL
    ws["E4"].alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells("I4:L4")
    ws["I4"] = "PACKING METHOD & MATERIAL FACTORS"
    ws["I4"].font = FONT_PARAM_HDR
    ws["I4"].fill = PARAM_HDR_FILL
    ws["I4"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[4].height = 22

    # Parameter grid setup (Rows 5 to 8)
    param_rows = [
        # Row 5
        [
            ("A5:B5", "40GP Internal Length (m)", PARAM_LBL_FILL, FONT_PARAM_LBL, BORDER_DATA, None),
            ("C5:D5", 12.032, CALC_CELL_FILL, FONT_PARAM_VAL, BORDER_DATA, "0.000"),
            ("E5:F5", "Selected Pipe Length (m) [Input]", PARAM_LBL_FILL, FONT_PARAM_LBL, BORDER_DATA, None),
            ("G5:H5", 11.80, INPUT_CELL_FILL, FONT_PARAM_INP, BORDER_INPUT, "0.00"),
            ("I5:J5", "Lengthwise Tiers in 40GP", PARAM_LBL_FILL, FONT_PARAM_LBL, BORDER_DATA, None),
            ("K5:L5", "=INT(C5/G5)", CALC_CELL_FILL, FONT_PARAM_VAL, BORDER_DATA, "0"),
        ],
        # Row 6
        [
            ("A6:B6", "40GP Internal Width (mm)", PARAM_LBL_FILL, FONT_PARAM_LBL, BORDER_DATA, None),
            ("C6:D6", 2352, CALC_CELL_FILL, FONT_PARAM_VAL, BORDER_DATA, "#,##0"),
            ("E6:F6", "Effective Loading Width (mm) [Input]", PARAM_LBL_FILL, FONT_PARAM_LBL, BORDER_DATA, None),
            ("G6:H6", 2300, INPUT_CELL_FILL, FONT_PARAM_INP, BORDER_INPUT, "#,##0"),
            ("I6:J6", "Cross-Section Stacking [Input]", PARAM_LBL_FILL, FONT_PARAM_LBL, BORDER_DATA, None),
            ("K6:L6", "Staggered (Hexagonal)", INPUT_CELL_FILL, FONT_PARAM_INP, BORDER_INPUT, None),
        ],
        # Row 7
        [
            ("A7:B7", "40GP Internal Height (mm)", PARAM_LBL_FILL, FONT_PARAM_LBL, BORDER_DATA, None),
            ("C7:D7", 2393, CALC_CELL_FILL, FONT_PARAM_VAL, BORDER_DATA, "#,##0"),
            ("E7:F7", "Effective Loading Height (mm) [Input]", PARAM_LBL_FILL, FONT_PARAM_LBL, BORDER_DATA, None),
            ("G7:H7", 2280, INPUT_CELL_FILL, FONT_PARAM_INP, BORDER_INPUT, "#,##0"),
            ("I7:J7", "Carbon Steel Factor C (kg/m)", PARAM_LBL_FILL, FONT_PARAM_LBL, BORDER_DATA, None),
            ("K7:L7", "='3_ASME_Spec_DB'!D6", CALC_CELL_FILL, FONT_PARAM_VAL, BORDER_DATA, "0.0000000"),
        ],
        # Row 8
        [
            ("A8:B8", "Max Container Payload (MT) [Input]", PARAM_LBL_FILL, FONT_PARAM_LBL, BORDER_DATA, None),
            ("C8:D8", 26.00, INPUT_CELL_FILL, FONT_PARAM_INP, BORDER_INPUT, "0.00"),
            ("E8:F8", "Max Container Payload (kg)", PARAM_LBL_FILL, FONT_PARAM_LBL, BORDER_DATA, None),
            ("G8:H8", "=C8*1000", CALC_CELL_FILL, FONT_PARAM_VAL, BORDER_DATA, "#,##0"),
            ("I8:J8", "SS304 / SS316 Factor C (kg/m)", PARAM_LBL_FILL, FONT_PARAM_LBL, BORDER_DATA, None),
            ("K8", "='3_ASME_Spec_DB'!D7", CALC_CELL_FILL, FONT_PARAM_VAL, BORDER_DATA, "0.0000000"),
            ("L8", "='3_ASME_Spec_DB'!D8", CALC_CELL_FILL, FONT_PARAM_VAL, BORDER_DATA, "0.0000000"),
        ],
    ]

    for r_idx, row_items in enumerate(param_rows, start=5):
        ws.row_dimensions[r_idx].height = 20
        for rng, val, fill, font, border, num_fmt in row_items:
            if ":" in rng:
                ws.merge_cells(rng)
                top_left_coord = rng.split(":")[0]
                cell_coords = rng.split(":")
            else:
                top_left_coord = rng
                cell_coords = [rng]
            cell = ws[top_left_coord]
            cell.value = val
            cell.font = font
            cell.fill = fill
            is_label = isinstance(val, str) and not val.startswith("=") and top_left_coord[0] in ("A", "E", "I")
            cell.alignment = Alignment(horizontal="left" if is_label else "center", vertical="center", indent=1 if is_label else 0)
            if num_fmt:
                cell.number_format = num_fmt
            for col_let in cell_coords:
                ws[col_let].border = border
                ws[col_let].fill = fill

    # Add Data Validations for Pipe Length (G5:H5) and Stacking Method (K6:L6)
    dv_len = DataValidation(type="list", formula1='"5.80,6.00,11.80,12.00"', allow_blank=False)
    ws.add_data_validation(dv_len)
    dv_len.add("G5:H5")

    dv_stack = DataValidation(type="list", formula1='"Staggered (Hexagonal),Square Grid"', allow_blank=False)
    ws.add_data_validation(dv_stack)
    dv_stack.add("K6:L6")

    # 3. Master Selectable Table Headers (Row 10)
    headers = [
        ("ID", TBL_HDR_MAIN),
        ("Material", TBL_HDR_MAIN),
        ("Standard", TBL_HDR_MAIN),
        ("NPS (in)", TBL_HDR_MAIN),
        ("DN (mm)", TBL_HDR_MAIN),
        ("Schedule", TBL_HDR_MAIN),
        ("OD (in)", TBL_HDR_MAIN),
        ("OD (mm)", TBL_HDR_MAIN),
        ("WT (in)", TBL_HDR_MAIN),
        ("WT (mm)", TBL_HDR_MAIN),
        ("Unit Weight\n(kg/m)", TBL_HDR_CALC),
        ("Piece Weight\n(kg/pc)", TBL_HDR_CALC),
        ("Horiz. Fit\n(pcs/row)", TBL_HDR_CALC),
        ("Vert. Layers\n(rows)", TBL_HDR_CALC),
        ("Max Pieces\nby Volume", TBL_HDR_CALC),
        ("Max Pieces\nby Weight", TBL_HDR_CALC),
        ("40GP Actual\nPieces (pcs)", TBL_HDR_OUT),
        ("40GP Actual\nMetric Tons (MT)", TBL_HDR_OUT),
        ("Payload\nUtil. (%)", TBL_HDR_CALC),
        ("Space Slot\nUtil. (%)", TBL_HDR_CALC),
        ("Limiting Factor\n(Bottleneck)", TBL_HDR_OUT),
    ]

    hdr_row = 10
    ws.row_dimensions[hdr_row].height = 30
    for c_idx, (h_text, h_fill) in enumerate(headers, start=1):
        c = ws.cell(row=hdr_row, column=c_idx, value=h_text)
        c.font = FONT_TBL_HDR
        c.fill = h_fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER_HDR

    # Build all valid rows for Carbon Steel, Stainless Steel 304, and Stainless Steel 316
    master_records = []
    # 1) Carbon Steel records (ASME B36.10M)
    for nps_str, dn_val, od_in, sch_dict in ASME_PIPE_SPECS:
        for sch in CARBON_SCHEDULES:
            if sch in sch_dict:
                master_records.append((
                    "Carbon Steel",
                    "ASME B36.10M",
                    nps_str,
                    dn_val,
                    sch,
                    od_in,
                    sch_dict[sch]
                ))

    # 2) Stainless Steel 304 records (ASME B36.19M)
    for nps_str, dn_val, od_in, sch_dict in ASME_PIPE_SPECS:
        for sch in STAINLESS_SCHEDULES:
            if sch in sch_dict:
                master_records.append((
                    "Stainless Steel 304",
                    "ASME B36.19M",
                    nps_str,
                    dn_val,
                    sch,
                    od_in,
                    sch_dict[sch]
                ))

    # 3) Stainless Steel 316 records (ASME B36.19M)
    for nps_str, dn_val, od_in, sch_dict in ASME_PIPE_SPECS:
        for sch in STAINLESS_SCHEDULES:
            if sch in sch_dict:
                master_records.append((
                    "Stainless Steel 316",
                    "ASME B36.19M",
                    nps_str,
                    dn_val,
                    sch,
                    od_in,
                    sch_dict[sch]
                ))

    start_data_row = 11
    for idx, (mat_name, std_name, nps_str, dn_val, sch_code, od_in, wt_in) in enumerate(master_records):
        r = start_data_row + idx
        ws.row_dimensions[r].height = 19
        row_fill = ZEBRA_FILL if idx % 2 == 1 else WHITE_FILL

        od_mm = round(od_in * INCH_TO_MM, 2)
        wt_mm = round(wt_in * INCH_TO_MM, 2)

        # Dynamic Excel Formulas for row r:
        # Col K (11): Unit Weight (kg/m) = (OD_mm - WT_mm) * WT_mm * Factor_C
        f_unit_wt = f'=ROUND((H{r}-J{r})*J{r}*IF(B{r}="Carbon Steel",$K$7,IF(B{r}="Stainless Steel 304",$K$8,$L$8)), 3)'
        # Col L (12): Piece Weight (kg/pc) = Unit_Weight * Pipe_Length
        f_pc_wt = f'=ROUND(K{r}*$G$5, 2)'
        # Col M (13): Horizontal Fit (bottom row count) = INT(Eff_Width / OD_mm)
        f_horiz = f'=INT($G$6/H{r})'
        # Col N (14): Vertical Layers (rows)
        f_vert = (
            f'=IF(M{r}=0, 0, IF(OR($K$6="Square Grid", M{r}=1), INT($G$7/H{r}), '
            f'IF($G$7<H{r}, 0, 1+INT(($G$7-H{r})/(0.8660254*H{r})))))'
        )
        # Col O (15): Max Pieces by Volume = Cross-Section Count * Lengthwise Tiers ($K$5)
        f_max_vol = (
            f'=IF(OR(M{r}=0,N{r}=0), 0, IF(OR($K$6="Square Grid", M{r}=1), M{r}*N{r}*$K$5, '
            f'(INT((N{r}+1)/2)*M{r} + INT(N{r}/2)*IF(($G$6-M{r}*H{r})>=0.5*H{r}, M{r}, M{r}-1))*$K$5))'
        )
        # Col P (16): Max Pieces by Weight = INT(Max_Payload_kg / Piece_Weight_kg)
        f_max_wt = f'=IF(L{r}<=0, 0, INT($G$8/L{r}))'
        # Col Q (17): 40GP Actual Pieces = MIN(Max_Vol, Max_Wt)
        f_act_pcs = f'=MIN(O{r}, P{r})'
        # Col R (18): 40GP Actual Metric Tons (MT) = ROUND((Actual_Pieces * Piece_Weight_kg) / 1000, 2)
        f_act_mt = f'=ROUND((Q{r}*L{r})/1000, 2)'
        # Col S (19): Payload Utilization (%) = Actual_MT / Max_Payload_MT
        f_wt_util = f'=IF($C$8>0, R{r}/$C$8, 0)'
        # Col T (20): Space Slot Utilization (%) = Actual_Pieces / Max_Pieces_by_Volume
        f_vol_util = f'=IF(O{r}>0, Q{r}/O{r}, 0)'
        # Col U (21): Limiting Factor (Bottleneck)
        f_limit = f'=IF(Q{r}=0, "Exceeds Limit", IF(P{r}<=O{r}, "Weight Limited", "Volume Limited"))'

        row_values = [
            idx + 1,        # A: ID
            mat_name,       # B: Material
            std_name,       # C: Standard
            nps_str,        # D: NPS (in)
            dn_val,         # E: DN (mm)
            sch_code,       # F: Schedule
            od_in,          # G: OD (in)
            od_mm,          # H: OD (mm)
            wt_in,          # I: WT (in)
            wt_mm,          # J: WT (mm)
            f_unit_wt,      # K: Unit Weight (kg/m)
            f_pc_wt,        # L: Piece Weight (kg/pc)
            f_horiz,        # M: Horiz. Fit (pcs/row)
            f_vert,         # N: Vert. Layers (rows)
            f_max_vol,      # O: Max Pieces by Volume
            f_max_wt,       # P: Max Pieces by Weight
            f_act_pcs,      # Q: 40GP Actual Pieces (pcs)
            f_act_mt,       # R: 40GP Actual Metric Tons (MT)
            f_wt_util,      # S: Payload Util. (%)
            f_vol_util,     # T: Space Slot Util. (%)
            f_limit,        # U: Limiting Factor
        ]

        for c_idx, val in enumerate(row_values, start=1):
            cell = ws.cell(row=r, column=c_idx, value=val)
            cell.border = BORDER_DATA

            # Styling & alignment by column type
            if c_idx in (17, 18):
                cell.fill = OUTPUT_COL_FILL
                cell.font = FONT_OUTPUT
            else:
                cell.fill = row_fill
                if c_idx == 2:
                    cell.font = FONT_CS_TAG if mat_name == "Carbon Steel" else FONT_SS_TAG
                elif c_idx in (4, 6, 21):
                    cell.font = FONT_BOLD
                else:
                    cell.font = FONT_REGULAR

            # Alignment
            if c_idx in (1, 3, 4, 5, 6, 13, 14, 21):
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif c_idx == 2:
                cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
            else:
                cell.alignment = Alignment(horizontal="right", vertical="center")

            # Number formats
            if c_idx in (1, 5, 13, 14, 15, 16, 17):
                cell.number_format = "#,##0"
            elif c_idx in (7, 9, 11):
                cell.number_format = "0.000"
            elif c_idx in (8, 10, 12, 18):
                cell.number_format = "#,##0.00"
            elif c_idx in (19, 20):
                cell.number_format = "0.0%"

    last_row = start_data_row + len(master_records) - 1

    # Add Conditional Formatting for Limiting Factor (Col U)
    wt_lim_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
    wt_lim_font = Font(name="Calibri", size=10, bold=True, color="B45309")
    vol_lim_fill = PatternFill(start_color="DBEAFE", end_color="DBEAFE", fill_type="solid")
    vol_lim_font = Font(name="Calibri", size=10, bold=True, color="1D4ED8")

    ws.conditional_formatting.add(
        f"U11:U{last_row}",
        CellIsRule(operator="equal", formula=['"Weight Limited"'], fill=wt_lim_fill, font=wt_lim_font)
    )
    ws.conditional_formatting.add(
        f"U11:U{last_row}",
        CellIsRule(operator="equal", formula=['"Volume Limited"'], fill=vol_lim_fill, font=vol_lim_font)
    )

    # Enable AutoFilter on the entire Master Table (Row 10 to last_row)
    ws.auto_filter.ref = f"A10:U{last_row}"

    # Freeze Panes so Control Block & Table Headers stay visible while scrolling
    ws.freeze_panes = "G11"

    # Set clean column widths
    col_widths = {
        1: 7,   # ID
        2: 20,  # Material
        3: 15,  # Standard
        4: 11,  # NPS
        5: 10,  # DN
        6: 11,  # Schedule
        7: 10,  # OD (in)
        8: 11,  # OD (mm)
        9: 10,  # WT (in)
        10: 11, # WT (mm)
        11: 14, # Unit Weight
        12: 14, # Piece Weight
        13: 12, # Horiz Fit
        14: 12, # Vert Layers
        15: 14, # Max Vol
        16: 14, # Max Wt
        17: 16, # 40GP Actual Pieces
        18: 18, # 40GP Actual MT
        19: 13, # Payload Util
        20: 13, # Space Util
        21: 18, # Limiting Factor
    }
    for c_idx, width in col_widths.items():
        ws.column_dimensions[get_column_letter(c_idx)].width = width

    wb.save(filename)
    print(f"Successfully updated {filename} with 2_Selectable_Master_Table ({len(master_records)} total pipe combinations).")


if __name__ == "__main__":
    build_step2_workbook()
