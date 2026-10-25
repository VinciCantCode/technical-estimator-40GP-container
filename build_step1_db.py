"""
Step 1: Build the Engineering Specification Database Sheet (3_ASME_Spec_DB)
for the Precise Pipe 40GP Container Loading Model.
"""

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from pipe_data import ASME_PIPE_SPECS, SCHEDULES

INCH_TO_MM = 25.4


def create_step1_workbook(filename="Precise_Pipe_40GP_Container_Model.xlsx"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "3_ASME_Spec_DB"
    ws.views.sheetView[0].showGridLines = True

    # Color Palette (Precise Pipe Industrial Navy & Steel Theme)
    NAVY_HEADER = PatternFill(start_color="0F2C59", end_color="0F2C59", fill_type="solid")
    SLATE_SUBHEADER = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    SS_HEADER = PatternFill(start_color="0D9488", end_color="0D9488", fill_type="solid")
    CS_HEADER = PatternFill(start_color="334155", end_color="334155", fill_type="solid")
    ZEBRA_FILL = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    WHITE_FILL = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    NA_FILL = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    ACCENT_BG = PatternFill(start_color="EFF6FF", end_color="EFF6FF", fill_type="solid")

    FONT_TITLE = Font(name="Calibri", size=15, bold=True, color="FFFFFF")
    FONT_SUBTITLE = Font(name="Calibri", size=10, italic=True, color="E2E8F0")
    FONT_SEC_HDR = Font(name="Calibri", size=11, bold=True, color="0F2C59")
    FONT_TBL_HDR = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    FONT_BOLD = Font(name="Calibri", size=10, bold=True, color="0F172A")
    FONT_REGULAR = Font(name="Calibri", size=10, bold=False, color="1E293B")
    FONT_NA = Font(name="Calibri", size=9, italic=True, color="94A3B8")

    THIN_SIDE = Side(border_style="thin", color="CBD5E1")
    MED_NAVY = Side(border_style="medium", color="0F2C59")
    BORDER_DATA = Border(left=THIN_SIDE, right=THIN_SIDE, top=THIN_SIDE, bottom=THIN_SIDE)
    BORDER_HDR = Border(left=THIN_SIDE, right=THIN_SIDE, top=MED_NAVY, bottom=MED_NAVY)

    # 1. Title Banner
    ws.merge_cells("A1:V1")
    ws["A1"] = "PRECISE PIPE — ASME B36.10M (CARBON STEEL) & B36.19M (STAINLESS STEEL) ENGINEERING DATABASE"
    ws["A1"].font = FONT_TITLE
    ws["A1"].fill = NAVY_HEADER
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[1].height = 30

    ws.merge_cells("A2:V2")
    ws["A2"] = "Reference Standard: ASME B36.10M / B36.19M | Nominal Pipe Size (NPS 1/8\" to 48\") | Outer Diameter (OD) & Wall Thickness (WT) in Millimeters (mm)"
    ws["A2"].font = FONT_SUBTITLE
    ws["A2"].fill = SLATE_SUBHEADER
    ws["A2"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[2].height = 20

    # 2. Engineering Constants Reference Block (Rows 4 to 8)
    ws["A4"] = "ENGINEERING CONSTANTS & MATERIAL DENSITIES"
    ws["A4"].font = FONT_SEC_HDR

    constants_headers = ["Material Grade", "Standard Spec", "Density (g/cm3)", "Weight Factor C (kg/m)", "Formula: Unit Weight (kg/m)"]
    for col_idx, text in enumerate(constants_headers, start=1):
        cell = ws.cell(row=5, column=col_idx, value=text)
        cell.font = FONT_TBL_HDR
        cell.fill = CS_HEADER
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER_HDR
    ws.row_dimensions[5].height = 22

    materials_info = [
        ("Carbon Steel (CS)", "ASME B36.10M / ASTM A53, A106, API 5L", 7.85, "=ROUND(PI()*10^-6*(C6*1000), 7)", "(OD_mm - WT_mm) * WT_mm * D6"),
        ("Stainless Steel 304/304L (SS304)", "ASME B36.19M / ASTM A312 TP304", 7.93, "=ROUND(PI()*10^-6*(C7*1000), 7)", "(OD_mm - WT_mm) * WT_mm * D7"),
        ("Stainless Steel 316/316L (SS316)", "ASME B36.19M / ASTM A312 TP316", 7.98, "=ROUND(PI()*10^-6*(C8*1000), 7)", "(OD_mm - WT_mm) * WT_mm * D8"),
    ]
    for idx, (mat, spec, dens, factor, form_str) in enumerate(materials_info, start=6):
        ws.row_dimensions[idx].height = 20
        vals = [mat, spec, dens, factor, form_str]
        for c_i, v in enumerate(vals, start=1):
            c = ws.cell(row=idx, column=c_i, value=v)
            c.font = FONT_BOLD if c_i == 1 else FONT_REGULAR
            c.fill = ACCENT_BG
            c.border = BORDER_DATA
            c.alignment = Alignment(horizontal="left" if c_i in (1, 2, 5) else "right", vertical="center")
            if c_i == 3:
                c.number_format = "0.00"
            elif c_i == 4:
                c.number_format = "0.0000000"

    # 3. Table 1: Master Wall Thickness Matrix in Millimeters (mm) (Starts at Row 10)
    ws["A10"] = "MASTER PIPE SCHEDULE CHART — WALL THICKNESS IN MILLIMETERS (MM)"
    ws["A10"].font = FONT_SEC_HDR

    matrix_headers = ["NPS (in)", "DN (mm)", "OD (in)", "OD (mm)"] + [f"Sch {s}" for s in SCHEDULES]
    hdr_row = 11
    ws.row_dimensions[hdr_row].height = 26
    for col_idx, h_text in enumerate(matrix_headers, start=1):
        cell = ws.cell(row=hdr_row, column=col_idx, value=h_text)
        cell.font = FONT_TBL_HDR
        if col_idx <= 4:
            cell.fill = NAVY_HEADER
        elif h_text.endswith("S"):
            cell.fill = SS_HEADER
        else:
            cell.fill = CS_HEADER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER_HDR

    start_matrix_row = 12
    for idx, (nps_str, dn_val, od_in, sch_dict) in enumerate(ASME_PIPE_SPECS):
        r = start_matrix_row + idx
        ws.row_dimensions[r].height = 19
        row_fill = ZEBRA_FILL if idx % 2 == 1 else WHITE_FILL
        od_mm = round(od_in * INCH_TO_MM, 2)

        base_vals = [nps_str, dn_val, od_in, od_mm]
        for c_i, val in enumerate(base_vals, start=1):
            c = ws.cell(row=r, column=c_i, value=val)
            c.font = FONT_BOLD if c_i in (1, 4) else FONT_REGULAR
            c.fill = row_fill
            c.border = BORDER_DATA
            c.alignment = Alignment(horizontal="center" if c_i in (1, 2) else "right", vertical="center")
            if c_i == 3:
                c.number_format = "0.000"
            elif c_i == 4:
                c.number_format = "0.00"

        for s_idx, sch_code in enumerate(SCHEDULES):
            c_col = 5 + s_idx
            c = ws.cell(row=r, column=c_col)
            c.border = BORDER_DATA
            if sch_code in sch_dict:
                wt_mm = round(sch_dict[sch_code] * INCH_TO_MM, 2)
                c.value = wt_mm
                c.font = FONT_REGULAR
                c.fill = row_fill
                c.alignment = Alignment(horizontal="right", vertical="center")
                c.number_format = "0.00"
            else:
                c.value = "-"
                c.font = FONT_NA
                c.fill = NA_FILL
                c.alignment = Alignment(horizontal="center", vertical="center")

    # 4. Table 2: Master Wall Thickness Matrix in Inches (in) (Starts at Row 50)
    in_title_row = start_matrix_row + len(ASME_PIPE_SPECS) + 2
    ws.cell(row=in_title_row, column=1, value="MASTER PIPE SCHEDULE CHART — WALL THICKNESS IN INCHES (IN)").font = FONT_SEC_HDR
    in_hdr_row = in_title_row + 1
    ws.row_dimensions[in_hdr_row].height = 26

    for col_idx, h_text in enumerate(matrix_headers, start=1):
        cell = ws.cell(row=in_hdr_row, column=col_idx, value=h_text)
        cell.font = FONT_TBL_HDR
        if col_idx <= 4:
            cell.fill = NAVY_HEADER
        elif h_text.endswith("S"):
            cell.fill = SS_HEADER
        else:
            cell.fill = CS_HEADER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER_HDR

    start_in_row = in_hdr_row + 1
    for idx, (nps_str, dn_val, od_in, sch_dict) in enumerate(ASME_PIPE_SPECS):
        r = start_in_row + idx
        ws.row_dimensions[r].height = 19
        row_fill = ZEBRA_FILL if idx % 2 == 1 else WHITE_FILL
        od_mm = round(od_in * INCH_TO_MM, 2)

        base_vals = [nps_str, dn_val, od_in, od_mm]
        for c_i, val in enumerate(base_vals, start=1):
            c = ws.cell(row=r, column=c_i, value=val)
            c.font = FONT_BOLD if c_i in (1, 3) else FONT_REGULAR
            c.fill = row_fill
            c.border = BORDER_DATA
            c.alignment = Alignment(horizontal="center" if c_i in (1, 2) else "right", vertical="center")
            if c_i == 3:
                c.number_format = "0.000"
            elif c_i == 4:
                c.number_format = "0.00"

        for s_idx, sch_code in enumerate(SCHEDULES):
            c_col = 5 + s_idx
            c = ws.cell(row=r, column=c_col)
            c.border = BORDER_DATA
            if sch_code in sch_dict:
                wt_in = sch_dict[sch_code]
                c.value = wt_in
                c.font = FONT_REGULAR
                c.fill = row_fill
                c.alignment = Alignment(horizontal="right", vertical="center")
                c.number_format = "0.000"
            else:
                c.value = "-"
                c.font = FONT_NA
                c.fill = NA_FILL
                c.alignment = Alignment(horizontal="center", vertical="center")

    # Auto-fit column widths cleanly
    col_widths = {
        1: 14,
        2: 12,
        3: 12,
        4: 12,
    }
    for col_idx in range(1, len(matrix_headers) + 1):
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = col_widths.get(col_idx, 10)

    # Freeze panes below the main matrix header
    ws.freeze_panes = "E12"

    wb.save(filename)
    print(f"Successfully generated {filename} with 3_ASME_Spec_DB ({len(ASME_PIPE_SPECS)} NPS sizes x {len(SCHEDULES)} Schedules).")


if __name__ == "__main__":
    create_step1_workbook()
